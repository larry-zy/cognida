"""OOXML presentations: slide text, table data and embedded image evidence."""

import base64
from io import BytesIO
from pathlib import Path
from typing import Any
from zipfile import ZipFile

from .base import BaseParser, ParseResult


class PPTXParser(BaseParser):
    @property
    def supported_formats(self):
        return ["pptx"]

    async def parse(
        self,
        source: str | bytes,
        include_metadata: bool = True,
        extract_tables: bool = False,
        extract_images: bool = False,
        **kwargs: Any,
    ) -> ParseResult:
        try:
            from pptx import Presentation

            content = Path(source).read_bytes() if isinstance(source, str) else source
            if len(content) > 50 * 1024 * 1024:
                raise ValueError("PPTX exceeds 50 MiB")
            with ZipFile(BytesIO(content)) as archive:
                entries = archive.infolist()
                if (
                    len(entries) > 10_000
                    or sum(entry.file_size for entry in entries) > 100 * 1024 * 1024
                ):
                    raise ValueError("PPTX expanded content exceeds extraction limits")
            presentation = Presentation(BytesIO(content))
            if len(presentation.slides) > 1000:
                raise ValueError("PPTX exceeds 1000 slides")
            parts, tables, images = [], [], []
            image_bytes = 0
            for page, slide in enumerate(presentation.slides, start=1):
                slide_text = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        slide_text.append(shape.text_frame.text)
                    if shape.has_table:
                        rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                        slide_text.extend("\t".join(row) for row in rows)
                        if extract_tables:
                            tables.append({"page": page, "table_index": len(tables), "rows": rows})
                    if extract_images and hasattr(shape, "image"):
                        raw = shape.image.blob
                        image_bytes += len(raw)
                        if image_bytes > 20 * 1024 * 1024:
                            raise ValueError("Extracted PPTX images exceed 20 MiB")
                        width, height = shape.image.size
                        images.append(
                            {
                                "page": page,
                                "image_index": len(images),
                                "width": width,
                                "height": height,
                                "ext": shape.image.ext,
                                "data": base64.b64encode(raw).decode(),
                                "path": f"pptx://slide/{page}/image/{len(images)}",
                            }
                        )
                parts.append("\n".join(slide_text))
            text = "\n\n".join(parts)
            metadata = {
                "format": "pptx",
                "page_count": len(presentation.slides),
                "char_count": len(text),
                "properties": {
                    "title": presentation.core_properties.title or "",
                    "author": presentation.core_properties.author or "",
                },
            }
            return ParseResult(
                True, text, metadata if include_metadata else {}, tables=tables, images=images
            )
        except ImportError:
            return ParseResult(
                False, "", {}, error="PPTX support requires the document extra (python-pptx)"
            )
        except Exception as exc:
            return ParseResult(False, "", {}, error=str(exc))
