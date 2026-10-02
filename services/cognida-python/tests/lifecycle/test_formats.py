import io

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from services.document import parse_document


async def test_html_parser_offline_text_tables_images():
    result = await parse_document(
        b'<html><title>Example</title><script>hidden()</script><p>Hello</p><table><tr><td>A</td><td>B</td></tr></table><img src="https://example.com/a.png" alt="chart"></html>',
        format="html",
        extract_tables=True,
        extract_images=True,
    )
    assert result["success"], result["error"]
    assert "Hello" in result["text"] and "hidden" not in result["text"]
    assert result["tables"][0]["rows"] == [["A", "B"]]
    assert result["metadata"]["properties"]["title"] == "Example"
    assert result["images"][0]["path"] == "https://example.com/a.png"


async def test_pptx_extracts_real_slide_table_and_image():
    document = Presentation()
    slide = document.slides.add_slide(document.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, Inches(3), Inches(1)).text = "Slide title"
    table = slide.shapes.add_table(2, 2, 0, Inches(1), Inches(4), Inches(2)).table
    table.cell(0, 0).text = "Name"
    table.cell(1, 0).text = "Alice"
    image = io.BytesIO()
    Image.new("RGB", (4, 3), "blue").save(image, format="PNG")
    image.seek(0)
    slide.shapes.add_picture(image, 0, Inches(4))
    out = io.BytesIO()
    document.save(out)
    result = await parse_document(
        out.getvalue(), format="pptx", extract_tables=True, extract_images=True
    )
    assert result["success"], result["error"]
    assert "Slide title" in result["text"] and "Alice" in result["text"]
    assert result["metadata"]["page_count"] == 1
    assert result["tables"][0]["rows"][1][0] == "Alice"
    assert result["images"][0]["width"] == 4 and result["images"][0]["data"]


@pytest.mark.parametrize("format", ["ppt", "doc", "xls"])
async def test_legacy_office_explicit_conversion_error(format):
    result = await parse_document(bytes.fromhex("d0cf11e0a1b11ae1") + b"old office", format=format)
    assert not result["success"]
    assert "convert" in result["error"].lower()
