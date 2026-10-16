"""Real OOXML bytes: searchable text must include tables and non-active sheets."""
from io import BytesIO

import pytest

from services.document.parsers.docx_parser import DocxParser
from services.document.parsers.excel_parser import ExcelParser


@pytest.mark.asyncio
async def test_docx_tables_are_searchable_without_table_metadata():
    from docx import Document
    document = Document()
    document.add_paragraph("BEFORE_TABLE")
    table = document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [["code", "amount"], ["00123", "9007199254740993.123456789012345678"]]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    document.add_paragraph("AFTER_TABLE")
    output = BytesIO()
    document.save(output)
    result = await DocxParser().parse(output.getvalue(), extract_tables=False)
    assert result.success
    assert result.text.index("BEFORE_TABLE") < result.text.index("00123") < result.text.index("AFTER_TABLE")
    assert "9007199254740993.123456789012345678" in result.text
    assert result.tables is None


@pytest.mark.asyncio
async def test_xlsx_all_sheets_default_and_explicit_sheet_scope():
    from openpyxl import Workbook
    workbook = Workbook()
    workbook.active.title = "First"
    workbook.active.append(["first evidence", "00123"])
    workbook.create_sheet("Second").append(["second evidence", "9007199254740993.123456789012345678"])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    result = await ExcelParser().parse(output.getvalue(), extract_tables=True)
    assert result.success
    assert "first evidence" in result.text and "second evidence" in result.text
    assert "00123" in result.text and "9007199254740993.123456789012345678" in result.text
    assert [table["sheet_name"] for table in result.tables] == ["First", "Second"]
    assert [table["table_index"] for table in result.tables] == [0, 1]
    selected = await ExcelParser().parse(output.getvalue(), sheet_name="Second")
    assert selected.success and "second evidence" in selected.text
    assert "first evidence" not in selected.text


@pytest.mark.asyncio
async def test_xlsx_invalid_explicit_sheet_closes_workbook(monkeypatch):
    import openpyxl
    from unittest.mock import Mock

    workbook = openpyxl.Workbook()
    close = Mock(wraps=workbook.close)
    monkeypatch.setattr(workbook, "close", close)
    monkeypatch.setattr(openpyxl, "load_workbook", lambda *_args, **_kwargs: workbook)
    result = await ExcelParser().parse(b"unused fixture bytes", sheet_name="missing")
    assert not result.success
    assert "missing" in result.error
    close.assert_called_once()
