"""Generate synthetic real-format bytes for isolated upload acceptance (no mocks)."""
import json
import sys
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

out = Path(sys.argv[1])
marker = sys.argv[2]
out.mkdir(parents=True, exist_ok=True)
manifest = []
for extension in ["txt", "md", "pdf", "docx", "xlsx", "csv", "html", "pptx"]:
    token = f"{marker}_{extension.upper()}"
    lines = [token, "The North region recorded revenue of 100 units in January.", "The South region recorded revenue of 200 units in February.", "The West region recorded revenue of 300 units in March.", "These synthetic observations describe this supplied sample only."]
    text = "\n".join(lines)
    path = out / f"{token}.{extension}"
    expected = [token, "North", "South", "West"]
    if extension == "txt":
        path.write_text(text)
    elif extension == "md":
        path.write_text(f"# {token}\n\n" + "\n\n".join(lines[1:]))
    elif extension == "csv":
        import csv
        with path.open("w", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(["marker", "region", "revenue", "code"])
            for region, revenue in [("North", 100), ("South", 200), ("West", 300)]:
                writer.writerow([token, region, revenue, "00123"])
        expected.append("00123")
    elif extension == "html":
        path.write_text(f"<!doctype html><html><head><title>{token}</title><script>SHOULD_NOT_PARSE</script></head><body>" + "".join(f"<p>{line}</p>" for line in lines) + "</body></html>")
    elif extension == "docx":
        document = Document()
        for line in lines:
            document.add_paragraph(line)
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "DOCX_TABLE_REQUIRED"
        table.cell(0, 1).text = "00123"
        document.add_paragraph("AFTER_REQUIRED_TABLE")
        document.save(path)
        expected.extend(["DOCX_TABLE_REQUIRED", "00123", "AFTER_REQUIRED_TABLE"])
    elif extension == "xlsx":
        workbook = Workbook()
        for line in lines:
            workbook.active.append([line])
        workbook.create_sheet("Additional evidence").append(["SECOND_SHEET_REQUIRED", "00123", "9007199254740993.123456789012345678"])
        workbook.save(path)
        workbook.close()
        expected.extend(["SECOND_SHEET_REQUIRED", "00123", "9007199254740993.123456789012345678"])
    elif extension == "pptx":
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        textbox = slide.shapes.add_textbox(Inches(.5), Inches(.4), Inches(9), Inches(4))
        textbox.text_frame.text = text
        table = slide.shapes.add_table(1, 2, Inches(.5), Inches(5), Inches(8), Inches(.5)).table
        table.cell(0, 0).text = "PPTX_TABLE_REQUIRED"
        table.cell(0, 1).text = "00123"
        presentation.save(path)
        expected.extend(["PPTX_TABLE_REQUIRED", "00123"])
    else:
        writer = PdfWriter()
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(("BT /F1 10 Tf 40 750 Td " + " ".join(f"({line}) Tj 0 -20 Td" for line in lines) + " ET").encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(stream)
        with path.open("wb") as handle:
            writer.write(handle)
    manifest.append({"format": extension, "path": str(path), "marker": token, "expected": expected})
# A real PDF without a text layer must fail honestly, not become a successful empty document.
blank = out / f"{marker}_NO_TEXT.pdf"
writer = PdfWriter()
writer.add_blank_page(width=612, height=792)
with blank.open("wb") as handle:
    writer.write(handle)
print(json.dumps({"files": manifest, "no_text_pdf": str(blank)}))
